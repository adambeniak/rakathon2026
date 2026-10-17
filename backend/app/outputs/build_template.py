"""Builds `templates/mdt_zapis.docx` from the sample form `data/raw/MDT_zapis.docx`.

Run: python -m app.outputs.build_template

The form's headings, tables and labels stay untouched; this only inserts docxtpl (Jinja) tags into the empty value
cells, a review-status line under each section heading, a draft banner under the title, and three appendices
(findings, excluded documents, remaining handover sections). Re-run it when the sample form changes and commit the
result; `tests/unit/outputs` renders the committed template.
"""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.settings import REPO_ROOT

SOURCE = REPO_ROOT / "data" / "raw" / "MDT_zapis.docx"
TARGET = Path(__file__).parent / "templates" / "mdt_zapis.docx"

HEADING_SIZE = Pt(13)


def _insert_after(anchor: Paragraph, text: str) -> Paragraph:
    element = OxmlElement("w:p")
    anchor._p.addnext(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.add_run(text)
    return paragraph


def _set_cell(table: Table, row: int, column: int, text: str) -> None:
    table.cell(row, column).paragraphs[0].text = text


def build(source: Path = SOURCE, target: Path = TARGET) -> Path:
    document = Document(str(source))
    headings = {p.text.split()[0]: p for p in document.paragraphs if p.text[:1].isdigit()}
    if sorted(headings) != ["1", "2", "3", "4", "5", "6", "7"]:
        raise ValueError("unexpected MDT_zapis layout: the seven numbered section headings are not all there")
    title = document.paragraphs[0]

    # Draft banner and the basis line under the title (inserted in reverse so the banner comes first).
    _insert_after(title, "{{ meta }}")
    _insert_after(title, "{{r banner }}")
    for number, section_heading in headings.items():
        _insert_after(section_heading, "{{r s" + number + " }}")

    tables = document.tables
    if len(tables) != 7:
        raise ValueError("unexpected MDT_zapis layout: expected seven tables")
    basics, _members, diagnosis, comorbidities, results, _conclusion, _recommendation = tables
    _set_cell(basics, 0, 1, "{{ mdt_date }}")
    _set_cell(basics, 1, 1, "{{ patient_lines }}")
    _set_cell(basics, 2, 1, "{{ referring_lines }}")
    _set_cell(diagnosis, 0, 1, "{{ diagnosis_name }}")
    _set_cell(diagnosis, 1, 1, "{{ diagnosis_code }}")
    _set_cell(comorbidities, 0, 0, "{{ comorbidities_text }}")
    # docxtpl row loop: the first and the last row hold only the tags and disappear when rendered.
    _set_cell(results, 1, 0, "{%tr for r in results %}")
    for column, tag in enumerate(("{{ r.date }}", "{{ r.exam }}", "{{ r.result }}")):
        _set_cell(results, 2, column, tag)
    _set_cell(results, 3, 0, "{%tr endfor %}")
    results._tbl.remove(results.rows[4]._tr)

    def add_heading(text: str) -> None:
        run = document.add_paragraph().add_run(text)
        run.bold = True
        run.font.size = HEADING_SIZE

    add_heading("Příloha A – Nesrovnalosti v dokumentaci (stav k MDT)")
    document.add_paragraph("{%p for line in findings %}")
    document.add_paragraph("– {{ line }}")
    document.add_paragraph("{%p endfor %}")
    add_heading("Příloha B – Dokumenty nezahrnuté do podkladu")
    document.add_paragraph("{%p for line in excluded %}")
    document.add_paragraph("– {{ line }}")
    document.add_paragraph("{%p endfor %}")
    add_heading("Příloha C – Další oddíly předávky (mimo formulář MDT zápisu)")
    document.add_paragraph("{%p for section in appendix %}")
    document.add_paragraph("{{r section.title }}")
    document.add_paragraph("{%p for line in section.lines %}")
    document.add_paragraph("– {{ line }}")
    document.add_paragraph("{%p endfor %}")
    document.add_paragraph("{%p endfor %}")
    footer = document.add_paragraph().add_run("{{ footer }}")
    footer.italic = True
    footer.font.size = Pt(8)

    # The draft banner repeats on every page (a printed page 3 or a forwarded file must still say "NÁVRH").
    header = document.sections[0].header
    header.is_linked_to_previous = False
    header.paragraphs[0].text = "{{r page_banner }}"

    _neutralise_metadata(document)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(target))
    return target


def _neutralise_metadata(document: DocumentType) -> None:
    """The sample form carries its author's name and a tenant's sensitivity-label ids; neither belongs in an export."""
    properties = document.core_properties
    fixed = datetime(2026, 10, 1, 0, 0, 0)  # noqa: DTZ001 - a fixed, timezone-less document timestamp
    properties.author = "Štafeta"
    properties.last_modified_by = "Štafeta"
    properties.title = "Zápis z jednání MDT"
    properties.subject = ""
    properties.comments = ""
    properties.keywords = ""
    properties.category = ""
    properties.created = fixed
    properties.modified = fixed
    package = document.part.package
    for rel_id, rel in list(package.rels.items()):
        if rel.reltype.endswith("/custom-properties"):
            del package.rels[rel_id]


if __name__ == "__main__":
    sys.stdout.write(f"{build()}\n")
