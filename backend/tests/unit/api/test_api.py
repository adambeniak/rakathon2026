from __future__ import annotations

import json
from collections.abc import Iterator
from datetime import date, datetime
from typing import Any

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.api.deps import (
    LlmClients,
    get_case_reader,
    get_directory_provider,
    get_fhir_client,
    get_handover_store,
    get_iris_pool,
    get_llm_clients,
)
from app.domain import clock
from app.domain.case import Case, CaseState
from app.domain.common import PartyRef, Role
from app.domain.document import SourceDocument
from app.domain.errors import NotFoundError
from app.domain.patient import Patient, Sex
from app.export_openapi import stable_json
from app.extraction.llm import CacheMissError
from app.fhir.cases import CaseContentsView, CaseDetailView, CaseOverview, LoadedCase
from app.fhir.directory import Directory
from app.fhir.errors import FhirAuthError, FhirConflictError, FhirUnavailableError
from app.main import create_app
from app.security.actor import current_actor
from app.settings import PRAGUE, Settings
from tests.directory_fixtures import demo_directory
from tests.unit.api.fake_handover_store import FakeStore

CASE = Case(
    id="rakathon-2026-prochazkova",
    patient=Patient(
        given="Alena",
        family="Procházková",
        sex=Sex.FEMALE,
        birth_date=date(1963, 7, 22),
        birth_number="635722/1013",
        insurer_code="211",
    ),
    state=CaseState.PRIPRAVENO_MDT,
    baton_holder=PartyRef(id="hruby-pavel@koc2-nemocnice-lipova", display="MDT GIT", role=Role.KOC_PHYSICIAN),
    next_deadline=datetime(2026, 10, 21, 8, 0, tzinfo=PRAGUE),
    opened_at=datetime(2026, 10, 2, 9, 30, tzinfo=PRAGUE),
)
DOCUMENT = SourceDocument(
    doc_id="07_PL_zadost_KOC2_2026-10-02",
    title="PL – žádost",
    type_code="57133-1",
    authored_at=date(2026, 10, 2),
    received_at=datetime(2026, 10, 2, 9, 30, tzinfo=PRAGUE),
)


class FakeReader:
    def __init__(self, error: Exception | None = None, mdt_start: datetime | None = None) -> None:
        self.error = error
        self.case = CASE.model_copy(update={"mdt_start": mdt_start})

    async def load(self, case_id: str) -> LoadedCase:
        if self.error:
            raise self.error
        if case_id != CASE.id:
            raise NotFoundError("case", case_id)
        return LoadedCase(task={"id": "5"}, episode={}, case=self.case)

    async def list_cases(self) -> list[CaseOverview]:
        if self.error:
            raise self.error
        return [CaseOverview(case=CASE, document_count=13)]

    async def get_case(self, case_id: str) -> CaseDetailView:
        if self.error:
            raise self.error
        if case_id != CASE.id:
            raise NotFoundError("case", case_id)
        return CaseDetailView(case=CASE, documents=[DOCUMENT])

    async def get_case_contents(self, case_id: str) -> CaseContentsView:
        detail = await self.get_case(case_id)
        documents = [*detail.documents, LATE_DOCUMENT]
        return CaseContentsView(case=CASE, documents=documents, contents={d.doc_id: _docx(d.title) for d in documents})


def _docx(paragraph: str) -> bytes:
    import io

    from docx import Document

    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph(paragraph)
    document.save(buffer)
    return buffer.getvalue()


class FakeLlm:
    """Extraction returns no facts; the epicrisis cache misses (-> cited fallback)."""

    def __init__(self, response: str | None) -> None:
        self.response = response
        self.calls: list[str] = []

    async def complete_json(self, *, system: str, user: str, schema: dict[str, object], doc_id: str) -> str:
        self.calls.append(doc_id)
        if self.response is None:
            raise CacheMissError("miss")
        return self.response


LATE_DOCUMENT = SourceDocument(
    doc_id="12_Algeziologie_2026-10-26",
    title="Algeziologie",
    type_code="11488-4",
    authored_at=date(2026, 10, 26),
    received_at=datetime(2026, 10, 26, 12, 0, tzinfo=PRAGUE),
)


class FakeFhir:
    def __init__(self, error: Exception | None = None) -> None:
        self.error = error

    async def capabilities(self) -> dict[str, str]:
        if self.error:
            raise self.error
        return {"resourceType": "CapabilityStatement"}


class FakePool:
    def __init__(self, ok: bool = True) -> None:
        self.ok = ok

    def ping(self) -> None:
        if not self.ok:
            raise ConnectionError("down")


@pytest.fixture
def app(settings: Settings) -> FastAPI:
    return create_app(settings)


@pytest.fixture
def client(app: FastAPI) -> Iterator[TestClient]:
    app.dependency_overrides[get_case_reader] = lambda: FakeReader()
    app.dependency_overrides[get_fhir_client] = lambda: FakeFhir()
    app.dependency_overrides[get_iris_pool] = lambda: FakePool()
    llm = LlmClients(FakeLlm('{"facts": []}'), FakeLlm(None))
    app.dependency_overrides[get_llm_clients] = lambda: llm
    store = FakeStore()
    app.dependency_overrides[get_handover_store] = lambda: store
    with TestClient(app) as test_client:
        yield test_client


def test_routers_are_auto_discovered(app: FastAPI) -> None:
    assert {"/api/health", "/api/cases", "/api/cases/{case_id}"} <= set(app.openapi()["paths"])


def test_health_ok(client: TestClient) -> None:
    response = client.get("/api/health")

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "ok"
    assert body["checks"]["fhir"]["ok"]
    assert body["checks"]["iris_sql"]["ok"]
    assert response.headers["X-Request-ID"]


def test_health_degraded_returns_503_with_reasons(app: FastAPI, client: TestClient) -> None:
    app.dependency_overrides[get_fhir_client] = lambda: FakeFhir(FhirAuthError("x", status_code=401))
    app.dependency_overrides[get_iris_pool] = lambda: FakePool(ok=False)

    response = client.get("/api/health")

    assert response.status_code == 503
    body = response.json()
    assert body["status"] == "degraded"
    assert body["checks"]["fhir"]["detail"] == "FhirAuthError (HTTP 401)"
    assert body["checks"]["iris_sql"]["detail"] == "ConnectionError"


def test_list_cases(client: TestClient) -> None:
    response = client.get("/api/cases")

    assert response.status_code == 200
    [case] = response.json()
    assert case["id"] == "rakathon-2026-prochazkova"
    assert case["patient"]["display_name"] == "Alena Procházková"
    assert case["patient"]["age_years"] == 63
    assert case["state"] == "pripraveno_mdt"
    assert case["baton_holder"]["role"] == "koc_physician"
    assert case["document_count"] == 13
    assert case["missing_count"] == case["counts"]["missing"]  # computed from the handover (test_case_counts.py)
    assert case["conflict_count"] == case["counts"]["conflict"]
    assert "birth_number" not in case["patient"]


def test_case_detail_lists_documents(client: TestClient) -> None:
    response = client.get("/api/cases/rakathon-2026-prochazkova")

    assert response.status_code == 200
    body = response.json()
    assert body["document_count"] == 1
    assert body["documents"][0]["doc_id"] == "07_PL_zadost_KOC2_2026-10-02"
    assert body["documents"][0]["received_at"] == "2026-10-02T09:30:00+02:00"


def test_unknown_case_is_problem_json(client: TestClient) -> None:
    response = client.get("/api/cases/unknown-case", headers={"X-Request-ID": "req-123"})

    assert response.status_code == 404
    assert response.headers["content-type"] == "application/problem+json"
    body = response.json()
    assert body["type"] == "https://stafeta.example/problems/not-found"
    assert body["request_id"] == "req-123"


def test_invalid_case_id_is_422_without_echoing_input(client: TestClient) -> None:
    response = client.get("/api/cases/bad id with spaces")

    assert response.status_code == 422
    body = response.json()
    assert body["type"] == "https://stafeta.example/problems/validation"
    assert body["errors"][0]["loc"] == ["path", "case_id"]
    assert all("input" not in error for error in body["errors"])


def test_storage_unavailable_is_503(app: FastAPI, client: TestClient) -> None:
    app.dependency_overrides[get_case_reader] = lambda: FakeReader(FhirUnavailableError("GET Task: ConnectError"))

    response = client.get("/api/cases")

    assert response.status_code == 503
    assert response.json()["type"] == "https://stafeta.example/problems/storage-unavailable"


def test_unhandled_error_is_500_problem_without_details(app: FastAPI, client: TestClient) -> None:
    app.dependency_overrides[get_case_reader] = lambda: FakeReader(RuntimeError("Alena Procházková 635722/1013"))

    response = client.get("/api/cases")

    assert response.status_code == 500
    assert response.headers["content-type"] == "application/problem+json"
    assert "Alena" not in response.text
    assert "Traceback" not in response.text


def test_openapi_is_stable_and_has_operation_ids(settings: Settings) -> None:
    first = stable_json(create_app(settings).openapi())
    second = stable_json(create_app(settings).openapi())

    assert first == second
    assert '"operationId": "cases-list_cases"' in first
    assert '"operationId": "cases-get_case"' in first
    assert '"operationId": "health-health"' in first


def test_access_log_has_full_route_template_and_no_ids(settings: Settings, capsys: pytest.CaptureFixture[str]) -> None:
    app = create_app(settings.model_copy(update={"log_level": "INFO"}))
    app.dependency_overrides[get_case_reader] = lambda: FakeReader()
    with TestClient(app) as test_client:
        test_client.get("/api/cases/rakathon-2026-prochazkova")

    lines = [json.loads(line) for line in capsys.readouterr().out.splitlines() if '"http_request"' in line]
    assert lines[-1]["route"] == "/api/cases/{case_id}"
    assert "rakathon-2026-prochazkova" not in json.dumps(lines[-1])


def test_405_problem_keeps_allow_header(client: TestClient) -> None:
    response = client.post("/api/cases")

    assert response.status_code == 405
    assert response.headers["content-type"] == "application/problem+json"
    assert "GET" in response.headers["allow"]


def test_openapi_declares_problem_json_for_errors(app: FastAPI) -> None:
    responses = app.openapi()["paths"]["/api/cases/{case_id}"]["get"]["responses"]

    assert list(responses["404"]["content"]) == ["application/problem+json"]
    assert list(responses["422"]["content"]) == ["application/problem+json"]
    assert list(responses["200"]["content"]) == ["application/json"]


def test_unhandled_error_response_carries_cors_headers(app: FastAPI, client: TestClient) -> None:
    app.dependency_overrides[get_case_reader] = lambda: FakeReader(RuntimeError("boom"))

    response = client.get("/api/cases", headers={"Origin": "http://localhost:5173"})

    assert response.status_code == 500
    assert response.headers["access-control-allow-origin"] == "http://localhost:5173"
    assert response.headers["X-Request-ID"]
    assert response.json()["request_id"] == response.headers["X-Request-ID"]


def test_cors_preflight_gets_a_request_id(client: TestClient) -> None:
    response = client.options(
        "/api/cases", headers={"Origin": "http://localhost:5173", "Access-Control-Request-Method": "GET"}
    )

    assert response.status_code == 200
    assert response.headers["X-Request-ID"]


def test_case_summary_has_the_header_fields_and_the_handover_counters(client: TestClient) -> None:
    [case] = client.get("/api/cases").json()

    assert case["postponed"] is False
    assert case["bucket"] == "aktivni"
    assert set(case["counts"]) == {"missing", "conflict", "stale"}  # values: tests/unit/api/test_case_counts.py
    assert case["suspected_diagnosis"] is None
    assert case["koc"] is None
    assert case["referring_gp"] is None


def test_handover_as_of_excludes_later_documents(client: TestClient) -> None:
    response = client.get("/api/cases/rakathon-2026-prochazkova/handover", params={"as_of": "2026-10-21"})

    assert response.status_code == 200
    body = response.json()
    assert body["as_of"] == "2026-10-21T08:00:00+02:00"
    assert body["included_documents"] == ["07_PL_zadost_KOC2_2026-10-02"]
    assert [e["doc_id"] for e in body["excluded_documents"]] == ["12_Algeziologie_2026-10-26"]
    fields = body["fields"]
    assert list(fields) == [f"{n}_" + k.split("_", 1)[1] for n, k in enumerate(fields, start=1)]
    assert fields["9_performance_status"]["status"] == "missing"
    assert fields["9_performance_status"]["missing_items"][0]["suggested_owner"] == "koc_physician"
    assert body["rules"]["stale_vitals_days"].startswith("30")


def test_handover_later_as_of_includes_documents(client: TestClient) -> None:
    response = client.get(
        "/api/cases/rakathon-2026-prochazkova/handover", params={"as_of": "2026-10-28T08:00:00+01:00"}
    )

    assert response.status_code == 200
    assert response.json()["excluded_documents"] == []
    assert "12_Algeziologie_2026-10-26" in response.json()["included_documents"]


@pytest.mark.parametrize("as_of", ["21. 10. 2026", "2026-10-21T08:00:00"])
def test_handover_invalid_as_of_is_422(client: TestClient, as_of: str) -> None:
    response = client.get("/api/cases/rakathon-2026-prochazkova/handover", params={"as_of": as_of})

    assert response.status_code == 422
    assert response.headers["content-type"] == "application/problem+json"


def test_handover_unknown_case_is_404(client: TestClient) -> None:
    response = client.get("/api/cases/unknown-case/handover")

    assert response.status_code == 404


def test_handover_export_markdown(client: TestClient) -> None:
    response = client.get(
        "/api/cases/rakathon-2026-prochazkova/handover/export", params={"as_of": "2026-10-21", "format": "md"}
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/markdown")
    assert response.text.startswith("# Předávka k MDT – stav k 21. 10. 2026 08:00")
    assert "## 9. PS / Karnofsky — chybí" in response.text
    assert "Nezahrnuto: 12_Algeziologie_2026-10-26" in response.text


class DemoDirectory:
    async def get(self) -> Directory:
        return demo_directory()


PHYSICIAN = PartyRef(id="markova-lenka@koc2-nemocnice-lipova", display="MUDr. Lenka Marková", role=Role.KOC_PHYSICIAN)
CONFIRM = "/api/cases/rakathon-2026-prochazkova/handover/sections/{n}/confirm?as_of=2026-10-21"


@pytest.fixture
def physician(app: FastAPI) -> None:
    app.dependency_overrides[current_actor] = lambda: PHYSICIAN


def _hash(client: TestClient, n: int) -> str:
    """value_hash of section n as the UI shows it (GET)."""
    fields = client.get("/api/cases/rakathon-2026-prochazkova/handover", params={"as_of": "2026-10-21"}).json()[
        "fields"
    ]
    value_hash: str = list(fields.values())[n - 1]["value_hash"]
    return value_hash


def _confirm(client: TestClient, n: int, headers: dict[str, str] | None = None, **body: object) -> Any:
    return client.post(CONFIRM.format(n=n), json={"expected_value_hash": _hash(client, n), **body}, headers=headers)


def test_persist_handover_returns_preliminary_composition(client: TestClient) -> None:
    response = client.post("/api/cases/rakathon-2026-prochazkova/handover", params={"as_of": "2026-10-21"})

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "preliminary"
    assert body["composition_id"] == "rakathon-2026-prochazkova@20261021T0600Z"
    assert body["fields"]["8_allergies"]["review"]["state"] == "draft"
    # GET now carries the stored review state too
    again = client.get("/api/cases/rakathon-2026-prochazkova/handover", params={"as_of": "2026-10-21"}).json()
    assert again["composition_id"] == body["composition_id"]


def test_confirm_requires_a_koc_physician(app: FastAPI, client: TestClient) -> None:
    # the real demo login (X-Demo-Role -> seeded person of that role), no actor override
    app.dependency_overrides[get_directory_provider] = lambda: DemoDirectory()

    response = _confirm(client, 9, headers={"X-Demo-Role": "coordinator"}, accept_missing=True)

    assert response.status_code == 403
    assert response.headers["content-type"] == "application/problem+json"
    assert response.json()["type"] == "https://stafeta.example/problems/forbidden"

    physician = _confirm(client, 9, headers={"X-Demo-Role": "koc_physician"}, accept_missing=True)
    assert physician.status_code == 200
    review = physician.json()["fields"]["9_performance_status"]["review"]
    assert review["by"]["role"] == "koc_physician"


@pytest.mark.usefixtures("physician")
def test_confirm_missing_section_needs_accept_missing(client: TestClient) -> None:
    response = _confirm(client, 9)

    assert response.status_code == 409
    assert response.json()["type"] == "https://stafeta.example/problems/conflict"


@pytest.mark.usefixtures("physician")
def test_accept_missing_records_who_and_when(client: TestClient) -> None:
    response = _confirm(client, 9, accept_missing=True)

    assert response.status_code == 200
    review = response.json()["fields"]["9_performance_status"]["review"]
    assert review["state"] == "accepted_missing"
    assert review["by"]["id"] == PHYSICIAN.id
    assert review["at"] == "2026-10-21T08:00:00+02:00"


@pytest.mark.usefixtures("physician")
def test_edited_value_is_validated_and_applied(client: TestClient) -> None:
    bad = _confirm(client, 9, edited_value={"scale": "WHO", "value": "1"})
    assert bad.status_code == 422
    assert "WHO" not in bad.text  # the submitted value is never echoed

    good = _confirm(client, 9, edited_value={"scale": "ECOG", "value": "1"})
    assert good.status_code == 200
    field = good.json()["fields"]["9_performance_status"]
    assert field["value"] == {"scale": "ECOG", "value": "1"}
    assert field["status"] == "complete"
    assert field["review"]["edited"] is True
    assert field["items"][0]["label"] == "Upraveno lékařem"


@pytest.mark.usefixtures("physician")
def test_section_number_out_of_range_is_422(client: TestClient) -> None:
    assert client.post(CONFIRM.format(n=16), json={"expected_value_hash": "0" * 64}).status_code == 422


@pytest.mark.usefixtures("physician")
def test_composition_final_after_all_sections(client: TestClient) -> None:
    body: dict[str, object] = {}
    for n in range(1, 16):
        response = _confirm(client, n, accept_missing=True)
        assert response.status_code == 200, (n, response.json())
        body = response.json()
    assert body["status"] == "final"


@pytest.mark.usefixtures("physician")
def test_confirm_with_outdated_hash_is_409(client: TestClient) -> None:
    response = client.post(CONFIRM.format(n=9), json={"expected_value_hash": "0" * 64, "accept_missing": True})

    assert response.status_code == 409
    assert "se od zobrazení změnil" in response.json()["detail"]


@pytest.mark.usefixtures("physician")
def test_confirm_without_hash_is_422(client: TestClient) -> None:
    assert client.post(CONFIRM.format(n=9), json={"accept_missing": True}).status_code == 422


def test_storage_version_conflict_is_409(app: FastAPI, client: TestClient) -> None:
    app.dependency_overrides[get_case_reader] = lambda: FakeReader(FhirConflictError("412", status_code=412))

    response = client.get("/api/cases")

    assert response.status_code == 409
    assert response.json()["type"] == "https://stafeta.example/problems/conflict"


MDT_START = datetime(2026, 10, 21, 8, 0, tzinfo=PRAGUE)


def test_handover_without_as_of_defaults_to_mdt_start(app: FastAPI, client: TestClient) -> None:
    app.dependency_overrides[get_case_reader] = lambda: FakeReader(mdt_start=MDT_START)

    with clock.frozen("2026-10-19T10:00:00"):
        response = client.get("/api/cases/rakathon-2026-prochazkova/handover")

    assert response.status_code == 200
    body = response.json()
    assert body["as_of"] == "2026-10-21T08:00:00+02:00"
    assert [d["doc_id"] for d in body["excluded_documents"]] == [LATE_DOCUMENT.doc_id]


def test_confirmations_in_different_minutes_hit_the_same_composition(app: FastAPI, client: TestClient) -> None:
    app.dependency_overrides[get_case_reader] = lambda: FakeReader(mdt_start=MDT_START)
    url = "/api/cases/rakathon-2026-prochazkova/handover"

    with clock.frozen("2026-10-19T10:00:00"):
        first = client.post(url).json()
    with clock.frozen("2026-10-19T10:07:00"):
        second = client.post(url).json()

    assert first["composition_id"] == second["composition_id"] == "rakathon-2026-prochazkova@20261021T0600Z"


def test_handover_without_as_of_and_without_mdt_start_uses_the_demo_clock(client: TestClient) -> None:
    with clock.frozen("2026-10-19T10:00:00"):
        response = client.get("/api/cases/rakathon-2026-prochazkova/handover")

    assert response.json()["as_of"] == "2026-10-19T10:00:00+02:00"


@pytest.mark.usefixtures("physician")
def test_step_3_decision_carries_over_to_the_mdt_start_handover(app: FastAPI, client: TestClient) -> None:
    """#122: the physician's decision on an unchanged section at the step-3 clock is not lost at mdt_start."""
    url = "/api/cases/rakathon-2026-prochazkova/handover"
    step3 = {"as_of": "2026-10-20T14:00:00+02:00"}
    shown = client.get(url, params=step3).json()["fields"]["9_performance_status"]
    confirmed = client.post(
        f"{url}/sections/9/confirm",
        params=step3,
        json={"expected_value_hash": shown["value_hash"], "accept_missing": True},
    )
    assert confirmed.status_code == 200
    app.dependency_overrides[get_case_reader] = lambda: FakeReader(mdt_start=MDT_START)

    with clock.frozen("2026-10-21T07:30:00"):
        mdt = client.post(url).json()

    assert mdt["composition_id"] == "rakathon-2026-prochazkova@20261021T0600Z"
    review = mdt["fields"]["9_performance_status"]["review"]
    assert review["state"] == "accepted_missing"
    assert review["by"]["id"] == PHYSICIAN.id
    assert mdt["fields"]["8_allergies"]["review"]["state"] == "draft"
