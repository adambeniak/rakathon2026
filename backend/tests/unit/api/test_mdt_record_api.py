"""GET /api/cases/{id}/mdt-record: KOC team only, audited, as of the MDT start, DOCX or PDF, signed conclusion."""

from __future__ import annotations

import io
from collections.abc import Iterator
from datetime import datetime
from typing import ClassVar

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

from app.api.deps import (
    LlmClients,
    get_case_reader,
    get_case_workflow,
    get_directory_provider,
    get_fhir_client,
    get_handover_store,
    get_iris_pool,
    get_llm_clients,
    get_repository,
)
from app.domain.case import Case
from app.domain.common import PartyRef
from app.domain.errors import NotFoundError
from app.domain.pathway import RecordedWaiver
from app.fhir.cases import CaseActivity, LoadedCase
from app.fhir.client import Resource
from app.fhir.directory import Directory
from app.fhir.mappers.mdt_conclusion import MdtConclusionRecord
from app.main import create_app
from app.outputs.mdt_record import DOCX_MEDIA_TYPE
from app.outputs.mdt_record_pdf import PDF_MEDIA_TYPE
from app.seed import demo_case
from app.settings import PRAGUE, Settings
from tests.directory_fixtures import demo_directory
from tests.unit.api.fake_handover_store import FakeStore
from tests.unit.api.test_api import FakeFhir, FakeLlm, FakePool, FakeReader

URL = "/api/cases/rakathon-2026-prochazkova/mdt-record"
MDT_START = datetime(2026, 10, 21, 8, 0, tzinfo=PRAGUE)
DIRECTORY = demo_directory()
GP = DIRECTORY.role(demo_case.GP_ROLE)
KOC2 = DIRECTORY.org(demo_case.KOC2)
assert GP
assert KOC2


class CaseReader(FakeReader):
    """The test_api reader, but `load` knows the referring GP and the managing KOC (authorize_case is real)."""

    loads = 0
    waivers: tuple[RecordedWaiver, ...] = ()

    async def load(self, case_id: str) -> LoadedCase:
        type(self).loads += 1
        base = await super().load(case_id)
        task = {
            "id": "5",
            "requester": {"reference": GP.ref},  # type: ignore[union-attr]
            "for": {"reference": "Patient/2"},
            "focus": {"reference": "EpisodeOfCare/1"},
        }
        episode = {"managingOrganization": {"reference": KOC2.ref}}  # type: ignore[union-attr]
        return LoadedCase(task=task, episode=episode, case=base.case)

    async def activity(self, task: Resource, case: Case, *, waivers: bool = True) -> CaseActivity:
        refs = tuple(f"Provenance/{60 + i}" for i, _ in enumerate(self.waivers))
        return CaseActivity(readiness_waivers=self.waivers, references=refs)


class FakeRepository:
    def __init__(self) -> None:
        self.exports: list[tuple[list[str], str]] = []

    async def audit_derived_write(self, entities: list[str], *, description: str) -> None:
        self.exports.append((entities, description))


class FakeProvider:
    async def get(self) -> Directory:
        return DIRECTORY


SIGNED = MdtConclusionRecord(
    composition_id="77",
    text="ZÁVĚR MDT\nText závěru dle jednání.\n\nDOPORUČENÍ MDT\n1. Bod podle lékaře.",
    signer_id="markova-lenka@koc2-nemocnice-lipova",
    signer_display="MUDr. Lenka Marková",
    signed_at=datetime(2026, 10, 21, 9, 30, tzinfo=PRAGUE),
)


class FakeWorkflow:
    """`get_conclusion` as the closed-loop workflow answers it: the signed record, or 404 before the sign-off."""

    conclusion: ClassVar[MdtConclusionRecord | None] = None
    reads: ClassVar[list[str]] = []

    async def get_conclusion(self, case_id: str, *, actor: PartyRef, directory: Directory) -> MdtConclusionRecord:
        type(self).reads.append(case_id)
        if self.conclusion is None:
            raise NotFoundError("MDT conclusion", case_id)
        return self.conclusion


def _client(settings: Settings, reader: CaseReader, repository: FakeRepository) -> Iterator[TestClient]:
    app = create_app(settings)
    app.dependency_overrides[get_case_reader] = lambda: reader
    app.dependency_overrides[get_directory_provider] = lambda: FakeProvider()
    app.dependency_overrides[get_repository] = lambda: repository
    app.dependency_overrides[get_case_workflow] = lambda: FakeWorkflow()
    app.dependency_overrides[get_fhir_client] = lambda: FakeFhir()
    app.dependency_overrides[get_iris_pool] = lambda: FakePool()
    llm = LlmClients(FakeLlm('{"facts": []}'), FakeLlm(None))
    app.dependency_overrides[get_llm_clients] = lambda: llm
    store = FakeStore()
    app.dependency_overrides[get_handover_store] = lambda: store
    with TestClient(app) as client:
        yield client


@pytest.fixture
def repository() -> FakeRepository:
    CaseReader.loads = 0
    CaseReader.waivers = ()
    FakeWorkflow.conclusion = None
    FakeWorkflow.reads = []
    return FakeRepository()


@pytest.fixture
def client(settings: Settings, repository: FakeRepository) -> Iterator[TestClient]:
    yield from _client(settings, CaseReader(mdt_start=MDT_START), repository)


@pytest.fixture
def client_without_mdt_start(settings: Settings, repository: FakeRepository) -> Iterator[TestClient]:
    yield from _client(settings, CaseReader(mdt_start=None), repository)


@pytest.mark.parametrize("role", ["coordinator", "koc_physician", "koc_staff"])
def test_the_koc_team_downloads_the_record_as_a_docx(client: TestClient, repository: FakeRepository, role: str) -> None:
    response = client.get(URL, headers={"X-Demo-Role": role})

    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX_MEDIA_TYPE
    assert response.headers["content-disposition"] == 'attachment; filename="mdt-zapis-NAVRH-2026-10-21.docx"'
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert "rakathon" not in response.headers["content-disposition"]  # the case id embeds the surname
    paragraphs = [p.text for p in Document(io.BytesIO(response.content)).paragraphs]
    assert "Zápis k jednání multidisciplinárního týmu (MDT)" in paragraphs
    assert "Stav podkladů k 21. 10. 2026 08:00. Zahrnuto dokumentů: 1, nezahrnuto: 1." in paragraphs
    assert any(p.startswith("– 12_Algeziologie_2026-10-26 – po MDT") for p in paragraphs)  # authored after the MDT
    # every export is audited, naming the patient and the case like a read does
    assert repository.exports[-1] == (["Task/5", "Patient/2", "EpisodeOfCare/1"], "MDT record exported as DOCX")


@pytest.mark.parametrize(("role", "status"), [("gp", 403), ("reviewer", 403)])
def test_only_the_koc_team_may_export(client: TestClient, repository: FakeRepository, role: str, status: int) -> None:
    response = client.get(URL, headers={"X-Demo-Role": role})

    assert response.status_code == status
    assert response.headers["content-type"].startswith("application/problem+json")
    assert repository.exports == []
    assert CaseReader.loads == 0  # refused by role before any audited read of clinical content
    assert FakeWorkflow.reads == []  # nor a read of the signed conclusion


def test_the_record_is_always_as_of_the_mdt_start(client: TestClient) -> None:
    other_day = client.get(URL, params={"as_of": "2026-10-28"}, headers={"X-Demo-Role": "coordinator"})

    document = Document(io.BytesIO(other_day.content))
    assert "Stav podkladů k 21. 10. 2026 08:00." in document.paragraphs[2].text  # a client cut-off is ignored


def _pdf_text(data: bytes) -> str:
    """Text of every page with whitespace collapsed (PDF extraction breaks lines where the layout wrapped them)."""
    return " ".join(" ".join(page.extract_text() for page in PdfReader(io.BytesIO(data)).pages).split())


def test_the_record_downloads_as_a_pdf_with_the_same_content(client: TestClient, repository: FakeRepository) -> None:
    response = client.get(URL, params={"format": "pdf"}, headers={"X-Demo-Role": "coordinator"})

    assert response.status_code == 200
    assert response.headers["content-type"] == PDF_MEDIA_TYPE
    assert response.headers["content-disposition"] == 'attachment; filename="mdt-zapis-NAVRH-2026-10-21.pdf"'
    assert response.headers["cache-control"] == "no-store"
    assert response.content.startswith(b"%PDF-")
    text = _pdf_text(response.content)
    assert "Zápis k jednání multidisciplinárního týmu (MDT)" in text
    assert "Stav podkladů k 21. 10. 2026 08:00. Zahrnuto dokumentů: 1, nezahrnuto: 1." in text
    assert "12_Algeziologie_2026-10-26 – po MDT" in text
    assert repository.exports[-1] == (["Task/5", "Patient/2", "EpisodeOfCare/1"], "MDT record exported as PDF")


def test_an_unknown_format_is_refused(client: TestClient, repository: FakeRepository) -> None:
    response = client.get(URL, params={"format": "odt"}, headers={"X-Demo-Role": "coordinator"})

    assert response.status_code == 422
    assert repository.exports == []


@pytest.mark.parametrize("fmt", ["docx", "pdf"])
def test_the_signed_conclusion_is_printed_in_section_6_and_its_read_is_audited(
    client: TestClient, repository: FakeRepository, fmt: str
) -> None:
    FakeWorkflow.conclusion = SIGNED

    response = client.get(URL, params={"format": fmt}, headers={"X-Demo-Role": "koc_physician"})

    assert response.status_code == 200
    if fmt == "docx":
        document = Document(io.BytesIO(response.content))
        text = "\n".join([p.text for p in document.paragraphs] + [c.text for t in document.tables for c in t._cells])
        assert SIGNED.text in text  # as signed, line breaks included
    else:
        text = _pdf_text(response.content)
        assert " ".join(SIGNED.text.split()) in text
    who = "MUDr. Lenka Marková, 21. 10. 2026 09:30"
    assert f"Potvrzeno lékařem – závěr zapsal(a) a v systému potvrdil(a) {who}" in text
    assert f"Závěr MDT zapsal(a) a v systému potvrdil(a) {who} (oddíl 6; nejde o elektronický podpis)." in text
    assert "12_Algeziologie_2026-10-26 – po MDT" in text  # the handover stays as of the MDT start
    entities, description = repository.exports[-1]
    assert entities == ["Task/5", "Patient/2", "EpisodeOfCare/1", "Composition/77"]
    assert description == f"MDT record exported as {fmt.upper()}"


def test_a_conclusion_signed_before_the_mdt_start_is_not_printed(
    client: TestClient, repository: FakeRepository
) -> None:
    FakeWorkflow.conclusion = MdtConclusionRecord(
        composition_id="70",
        text="Závěr jiného jednání.",
        signer_id="markova-lenka@koc2-nemocnice-lipova",
        signer_display="MUDr. Lenka Marková",
        signed_at=datetime(2026, 10, 14, 9, 30, tzinfo=PRAGUE),
    )

    response = client.get(URL, params={"format": "pdf"}, headers={"X-Demo-Role": "coordinator"})

    text = _pdf_text(response.content)
    assert "Závěr jiného jednání." not in text
    assert "Vyplní lékař na jednání MDT. Systém závěr ani doporučení nevytváří." in text
    assert repository.exports[-1][0] == ["Task/5", "Patient/2", "EpisodeOfCare/1"]


def test_a_case_without_a_set_mdt_start_has_no_record(
    client_without_mdt_start: TestClient, repository: FakeRepository
) -> None:
    response = client_without_mdt_start.get(URL, headers={"X-Demo-Role": "coordinator"})

    assert response.status_code == 422
    assert repository.exports == []


def test_an_unknown_case_is_404(client: TestClient) -> None:
    response = client.get("/api/cases/unknown-case/mdt-record", headers={"X-Demo-Role": "coordinator"})

    assert response.status_code == 404


def test_a_waived_readiness_item_is_in_the_record_and_its_read_is_audited(
    client: TestClient, repository: FakeRepository
) -> None:
    signed_at = datetime(2026, 10, 21, 7, 45, tzinfo=PRAGUE)
    CaseReader.waivers = (
        RecordedWaiver(
            "performance_status", "PS doplní lékař osobně na MDT.", "m@koc2", "MUDr. Lenka Marková", signed_at
        ),
    )

    response = client.get(URL, headers={"X-Demo-Role": "koc_physician"})

    assert response.status_code == 200
    paragraphs = [p.text for p in Document(io.BytesIO(response.content)).paragraphs]
    assert (
        "– Prominuto lékařem (připravenost k MDT – PS / Karnofsky): MUDr. Lenka Marková, 21. 10. 2026 07:45, "
        "důvod: PS doplní lékař osobně na MDT." in paragraphs
    )
    entities, _ = repository.exports[-1]
    assert entities == ["Task/5", "Patient/2", "EpisodeOfCare/1", "Provenance/60"]
