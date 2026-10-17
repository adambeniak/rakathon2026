"""MDT record export: layout of MDT_zapis, status labels, as-of, minimisation, no system-written conclusion."""

from __future__ import annotations

import io
from datetime import datetime

import pytest
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.domain.common import PartyRef, Role
from app.domain.handover import Handover, HandoverFields, HandoverStatus, ReviewState, SectionReview
from app.extraction.handover.review import overlay_review
from app.outputs.build_template import SOURCE
from app.outputs.mdt_record import (
    BANNER_FINAL,
    CONCLUSION_NOTE,
    LABEL_ACCEPTED_MISSING,
    LABEL_CONFIRMED,
    LABEL_CONFIRMED_EDITED,
    LABEL_DRAFT,
    LABEL_MISSING,
    RECOMMENDATION_IN_CONCLUSION,
    ReferringDoctor,
    SignedConclusion,
    field_label,
    render_mdt_record,
)
from app.settings import PRAGUE

PHYSICIAN = PartyRef(id="markova-lenka@koc2-nemocnice-lipova", display="MUDr. Lenka Marková", role=Role.KOC_PHYSICIAN)
GP = ReferringDoctor(name="MUDr. Tomáš Řezníček", phone="+420 311 000 287")
CONFIRMED_AT = datetime(2026, 10, 20, 15, 0, tzinfo=PRAGUE)
EXPORTED_AT = datetime(2026, 10, 21, 7, 30, tzinfo=PRAGUE)


def render(handover: Handover, referring: ReferringDoctor | None = GP) -> bytes:
    return render_mdt_record(handover, referring, EXPORTED_AT)


def _lines(data: bytes) -> list[str]:
    """Every paragraph and table cell of the document in body order (cell text joined by ' | ')."""
    document = Document(io.BytesIO(data))
    lines: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.split("}")[1]
        if tag == "p":
            lines.append(Paragraph(child, document).text)
        elif tag == "tbl":
            lines += [" | ".join(c.text for c in row.cells) for row in Table(child, document).rows]
    return lines


def _text(data: bytes) -> str:
    return "\n".join(_lines(data))


def _confirmed(handover: Handover, *, edited: bool = False) -> Handover:
    sections: dict[str, SectionReview] = {}
    for name, info in HandoverFields.model_fields.items():
        field = getattr(handover.fields, name)
        state = ReviewState.ACCEPTED_MISSING if field.status is HandoverStatus.MISSING else ReviewState.CONFIRMED
        sections[info.alias or name] = SectionReview(
            state=state, by=PHYSICIAN, at=CONFIRMED_AT, edited=edited, value_hash=field.value_hash
        )
    return overlay_review(handover, composition_id="c@x", status="final", sections=sections, edited_values={})


def test_the_template_keeps_the_layout_of_the_sample_form(story_handover: Handover) -> None:
    sample = [p.text for p in Document(str(SOURCE)).paragraphs if p.text[:1].isdigit()]

    rendered = [line for line in _lines(render(story_handover)) if line[:1].isdigit()]

    assert [h for h in rendered if h in sample] == sample  # the seven headings, same text, same order
    assert [h.split()[0] for h in sample] == ["1", "2", "3", "4", "5", "6", "7"]


def test_basic_data_diagnosis_and_results_come_from_the_handover(story_handover: Handover) -> None:
    text = _text(render(story_handover))

    assert "21. 10. 2026 08:00" in text
    assert "Alena Procházková" in text
    assert "211 – Zdravotní pojišťovna ministerstva vnitra ČR" in text
    assert "MUDr. Tomáš Řezníček" in text
    assert "+420 311 000 287" in text
    assert "Název diagnózy | Duktální adenokarcinom pankreatu" in text
    assert "Kód diagnózy (MKN-10) | C25.0" in text
    assert "Paroxysmální fibrilace síní (I48.0)" in text
    assert "CA 19-9 312 kU/l (H)" in text  # a lab value with its flag, Czech decimal comma elsewhere (4,1)
    assert "perineurální invaze: ano" in text


def test_without_confirmation_it_is_a_draft_with_labelled_sections(story_handover: Handover) -> None:
    lines = _lines(render(story_handover))
    text = "\n".join(lines)

    assert any(line.startswith("NÁVRH – zápis obsahuje oddíly, které lékař dosud nepotvrdil") for line in lines)
    assert LABEL_DRAFT in text
    assert LABEL_CONFIRMED not in text
    # Missing performance status is shown as missing, never inferred.
    assert "PS / Karnofsky – Chybí" in text
    assert "Chybí: ECOG / Karnofsky – ECOG/Karnofsky v dokumentaci neuvedeno" in text


def test_confirmed_sections_are_labelled_and_the_banner_drops_the_draft_warning(story_handover: Handover) -> None:
    lines = _lines(render(_confirmed(story_handover)))
    text = "\n".join(lines)

    assert BANNER_FINAL in lines
    assert not any(line.startswith("NÁVRH") for line in lines)
    assert LABEL_CONFIRMED in text
    assert LABEL_ACCEPTED_MISSING in text  # the missing performance status was accepted as missing
    assert LABEL_DRAFT not in text


def test_a_section_edited_by_the_physician_says_so() -> None:
    from app.domain.handover import PatientField

    field = PatientField(
        status=HandoverStatus.COMPLETE,
        review=SectionReview(state=ReviewState.CONFIRMED, by=PHYSICIAN, at=CONFIRMED_AT, edited=True),
    )

    assert field_label(field) == LABEL_CONFIRMED_EDITED
    assert field_label(PatientField(status=HandoverStatus.MISSING)) == LABEL_MISSING
    assert field_label(PatientField(status=HandoverStatus.PARTIAL)) == LABEL_DRAFT


def test_reports_after_the_mdt_are_listed_as_excluded_and_never_used(story_handover: Handover) -> None:
    lines = _lines(render(story_handover))
    excluded = {"12_Algeziologie_2026-10-26", "13_Psychologie_KOC2_2026-10-27"}

    appendix_b = lines.index("Příloha B – Dokumenty nezahrnuté do podkladu")
    appendix_c = lines.index("Příloha C – Další oddíly předávky (mimo formulář MDT zápisu)")
    listed = lines[appendix_b + 1 : appendix_c]
    assert {doc for doc in excluded if any(line.startswith(f"– {doc} – po MDT") for line in listed)} == excluded
    outside_b = "\n".join(lines[:appendix_b] + lines[appendix_c:])
    assert not any(doc in outside_b for doc in excluded)  # no value, evidence or finding cites them
    assert "as_of" not in "\n".join(lines)  # technical wording does not leak into the document
    assert "Zahrnuto dokumentů: 11, nezahrnuto: 2." in lines[2]


def test_findings_are_listed_with_their_documents(story_handover: Handover) -> None:
    lines = _lines(render(story_handover))

    start = lines.index("Příloha A – Nesrovnalosti v dokumentaci (stav k MDT)")
    end = lines.index("Příloha B – Dokumenty nezahrnuté do podkladu")
    findings = lines[start + 1 : end]
    assert {f.split(" ")[1] for f in findings} == {"C1", "C2"}
    assert any("09_Biopsie_vykon_KOC2_2026-10-12" in f for f in findings)


def test_conclusion_and_recommendation_are_left_to_the_physicians(story_handover: Handover) -> None:
    lines = _lines(render(story_handover))

    for heading in ("6  Závěr MDT", "7  Doporučení MDT"):
        at = lines.index(heading)
        assert lines[at + 1] == "Vyplní lékař na jednání MDT. Systém závěr ani doporučení nevytváří."
        assert lines[at + 2] == ""  # the empty table cell stays empty


def test_after_the_sign_off_section_6_prints_the_signed_conclusion_and_7_points_to_it(
    story_handover: Handover,
) -> None:
    text = "ZÁVĚR MDT\nZávěr podle jednání, RČ 635722/1013.\n\nDOPORUČENÍ MDT\n1. Bod zapsaný lékařem."
    signed = SignedConclusion(
        text=text, signer_display=PHYSICIAN.display, signed_at=datetime(2026, 10, 21, 9, 30, tzinfo=PRAGUE)
    )

    lines = _lines(render_mdt_record(story_handover, GP, EXPORTED_AT, conclusion=signed))

    at = lines.index("6  Závěr MDT")
    who = "MUDr. Lenka Marková, 21. 10. 2026 09:30"
    assert lines[at + 1] == f"{LABEL_CONFIRMED} – závěr zapsal(a) a v systému potvrdil(a) {who}\n{CONCLUSION_NOTE}"
    assert lines[at + 2] == text.replace("635722/1013", "[RČ vynecháno]")  # as signed; birth number left out
    at = lines.index("7  Doporučení MDT")
    assert lines[at + 1] == RECOMMENDATION_IN_CONCLUSION  # never split out of the signed text
    assert lines[at + 2] == ""
    meta = f"Závěr MDT zapsal(a) a v systému potvrdil(a) {who} (oddíl 6; nejde o elektronický podpis)."
    assert any(line.endswith(meta) for line in lines)
    assert "Závěr MDT podepsán" not in "\n".join(lines)  # no electronic signature in the demo: never claimed


def test_minimisation_no_birth_number_in_the_record(story_handover: Handover) -> None:
    text = _text(render(story_handover))

    assert "635722" not in text
    assert "rakathon" not in text.lower()  # the case id embeds the surname


def test_special_characters_are_escaped_not_interpreted(story_handover: Handover) -> None:
    diagnosis = story_handover.fields.diagnosis
    assert diagnosis.value is not None
    hostile = diagnosis.model_copy(update={"value": diagnosis.value.model_copy(update={"text": 'A & <B> "c" {{x}}'})})
    handover = story_handover.model_copy(
        update={"fields": story_handover.fields.model_copy(update={"diagnosis": hostile})}
    )

    text = _text(render(handover))

    assert 'Název diagnózy | A & <B> "c" {{x}}' in text


@pytest.mark.parametrize("referring", [None, ReferringDoctor(name=None)])
def test_an_unknown_referring_doctor_is_shown_as_missing(
    story_handover: Handover, referring: ReferringDoctor | None
) -> None:
    text = _text(render(story_handover, referring))

    assert "Referující lékař:" in text
    assert "Telefon: chybí · e-mail: chybí" in text


def test_a_waived_readiness_item_is_printed_next_to_its_section(story_handover: Handover) -> None:
    from app.domain.pathway import RecordedWaiver

    at = datetime(2026, 10, 21, 7, 45, tzinfo=PRAGUE)
    waivers = [
        RecordedWaiver("performance_status", "PS doplní lékař osobně na MDT.", PHYSICIAN.id, PHYSICIAN.display, at),
        RecordedWaiver("histology", "Histologie doručena ústně, RČ 635722/1013.", PHYSICIAN.id, PHYSICIAN.display, at),
        RecordedWaiver("pathway_extra_item", "Položka bez oddílu.", PHYSICIAN.id, None, at),
    ]

    lines = _lines(render_mdt_record(story_handover, GP, EXPORTED_AT, waivers))
    text = "\n".join(lines)

    ps = next(i for i, line in enumerate(lines) if line.startswith("PS / Karnofsky – Chybí"))
    assert (
        "– Prominuto lékařem (připravenost k MDT – PS / Karnofsky): MUDr. Lenka Marková, 21. 10. 2026 07:45, "
        "důvod: PS doplní lékař osobně na MDT." in lines[ps : ps + 6]
    )
    section5 = lines[lines.index("5 Výsledky stávajících vyšetření") + 1]  # the label paragraph of section 5
    assert "Prominuto lékařem (připravenost k MDT – Histologie): MUDr. Lenka Marková, 21. 10. 2026 07:45" in section5
    assert "635722/1013" not in text  # the reason is sanitised like every other value
    assert "Další prominuté položky připravenosti k MDT" in text
    assert "(připravenost k MDT – položka „pathway_extra_item“): jméno neuvedeno, 21. 10. 2026 07:45" in text
    assert text.count("Prominuto lékařem") == 3  # each waiver once, nowhere else


def test_without_waivers_the_record_says_nothing_about_them(story_handover: Handover) -> None:
    text = _text(render(story_handover))

    assert "Prominuto" not in text
    assert "Další prominuté položky" not in text
