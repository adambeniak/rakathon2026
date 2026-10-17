"""MDT record as PDF: same content and labels as the DOCX, Czech glyphs, draft marker on every page, escaping."""

from __future__ import annotations

import io
from datetime import datetime

import pytest
from docx import Document
from pypdf import PdfReader

from app.domain.handover import Handover
from app.domain.pathway import RecordedWaiver
from app.outputs import mdt_record_pdf
from app.outputs.build_template import SOURCE
from app.outputs.mdt_record import (
    BANNER_DRAFT,
    BANNER_FINAL,
    FOR_PHYSICIANS,
    PAGE_BANNER_DRAFT,
    RECOMMENDATION_IN_CONCLUSION,
    REDACTED_BIRTH_NUMBER,
    TEMPLATE,
    MdtRecordError,
    SignedConclusion,
    build_context,
)
from app.outputs.mdt_record_pdf import BASICS_LABELS, HEADINGS, render_mdt_record_pdf
from app.settings import PRAGUE
from tests.unit.outputs.test_mdt_record import EXPORTED_AT, GP, _confirmed

SIGNED_AT = datetime(2026, 10, 21, 9, 30, tzinfo=PRAGUE)
CONCLUSION = SignedConclusion(
    text="ZÁVĚR MDT\nText závěru podle jednání, ř š č ž ů.\n\nDOPORUČENÍ MDT\n1. Bod zapsaný lékařem.",
    signer_display="MUDr. Lenka Marková",
    signed_at=SIGNED_AT,
)
WAIVER = RecordedWaiver(
    "performance_status",
    "PS doplní lékař osobně na MDT.",
    "markova-lenka@koc2",
    "MUDr. Lenka Marková",
    datetime(2026, 10, 21, 7, 45, tzinfo=PRAGUE),
)


def _pages(data: bytes) -> list[str]:
    return [" ".join(page.extract_text().split()) for page in PdfReader(io.BytesIO(data)).pages]


def _text(data: bytes) -> str:
    return " ".join(_pages(data))


def _norm(text: str) -> str:
    return " ".join(text.split())


def test_the_pdf_has_the_form_headings_in_order_with_czech_glyphs(story_handover: Handover) -> None:
    text = _text(render_mdt_record_pdf(story_handover, GP, EXPORTED_AT))

    positions = [text.index(h) for h in HEADINGS]
    assert positions == sorted(positions)
    assert "Zápis k jednání multidisciplinárního týmu (MDT)" in text
    assert "Příloha A – Nesrovnalosti v dokumentaci (stav k MDT)" in text
    assert "Alena Procházková" in text  # diacritics survive the embedded font
    assert "MUDr. Tomáš Řezníček" in text


def test_the_headings_and_labels_are_the_sample_forms(story_handover: Handover) -> None:
    form = [_norm(p.text) for p in Document(str(SOURCE)).paragraphs if p.text[:1].isdigit()]
    template = Document(str(TEMPLATE))

    assert list(HEADINGS) == form
    assert [template.tables[0].cell(i, 0).text for i in range(3)] == list(BASICS_LABELS)


def test_the_bundled_font_is_embedded(story_handover: Handover) -> None:
    data = render_mdt_record_pdf(story_handover, GP, EXPORTED_AT)

    assert b"/FontFile2" in data  # TrueType program embedded (subset), no reliance on the reader's fonts
    assert b"+DejaVuSans" in data  # subset tag + PostScript name of the bundled face


def test_every_value_of_the_document_model_is_in_the_pdf(story_handover: Handover) -> None:
    """Same content as the DOCX: both render the one model of `build_context`."""
    context = build_context(story_handover, GP, EXPORTED_AT, (WAIVER,), CONCLUSION)

    text = _text(render_mdt_record_pdf(story_handover, GP, EXPORTED_AT, (WAIVER,), CONCLUSION))

    values = [context["meta"], context["footer"], context["patient_lines"], context["diagnosis_name"]]
    values += [context[k].text for k in ("s1", "s2", "s3", "s4", "s5", "s6", "s7")]
    values += context["findings"] + context["excluded"]
    values += [f"{r['exam']} {r['result']}" for r in context["results"]]
    for appendix in context["appendix"]:
        values += [appendix["title"].text, *appendix["lines"]]
    missing = [v for v in values if _norm(v) not in text]
    assert missing == []


def test_a_draft_is_marked_on_every_page(story_handover: Handover) -> None:
    pages = _pages(render_mdt_record_pdf(story_handover, GP, EXPORTED_AT))

    assert len(pages) > 1
    assert all(PAGE_BANNER_DRAFT in page for page in pages)
    assert BANNER_DRAFT in pages[0]


def test_a_fully_confirmed_record_has_no_draft_marker(story_handover: Handover) -> None:
    pages = _pages(render_mdt_record_pdf(_confirmed(story_handover), GP, EXPORTED_AT))

    assert not any(PAGE_BANNER_DRAFT in page for page in pages)
    assert BANNER_FINAL in pages[0]


def test_a_waiver_is_printed_as_recorded(story_handover: Handover) -> None:
    text = _text(render_mdt_record_pdf(story_handover, GP, EXPORTED_AT, (WAIVER,)))

    assert (
        "Prominuto lékařem (připravenost k MDT – PS / Karnofsky): MUDr. Lenka Marková, 21. 10. 2026 07:45, "
        "důvod: PS doplní lékař osobně na MDT." in text
    )


def test_without_a_signed_conclusion_sections_6_and_7_are_left_for_the_physicians(story_handover: Handover) -> None:
    text = _text(render_mdt_record_pdf(story_handover, GP, EXPORTED_AT))

    assert text.count(FOR_PHYSICIANS) == 2
    assert "závěr zapsal(a)" not in text


def test_the_signed_conclusion_is_printed_as_signed_with_who_and_when(story_handover: Handover) -> None:
    text = _text(render_mdt_record_pdf(story_handover, GP, EXPORTED_AT, conclusion=CONCLUSION))

    section_6 = text[text.index(HEADINGS[5]) : text.index(HEADINGS[6])]
    section_7 = text[text.index(HEADINGS[6]) : text.index("Příloha A")]
    assert (
        "Potvrzeno lékařem – závěr zapsal(a) a v systému potvrdil(a) MUDr. Lenka Marková, 21. 10. 2026 09:30"
        in section_6
    )
    assert "Nejde o elektronický podpis." in section_6
    assert _norm(CONCLUSION.text) in section_6
    assert _norm(RECOMMENDATION_IN_CONCLUSION) in section_7
    assert "Bod zapsaný lékařem" not in section_7  # never split out of the signed text
    assert FOR_PHYSICIANS not in text


def test_text_is_escaped_and_sanitised(story_handover: Handover) -> None:
    hostile = SignedConclusion(
        text='Závěr <b>tučně</b> & <font color="red">x</font>\fRČ 635722/1013 \x00konec',
        signer_display="MUDr. <i>Test</i>",
        signed_at=SIGNED_AT,
    )

    text = _text(render_mdt_record_pdf(story_handover, GP, EXPORTED_AT, conclusion=hostile))

    assert 'Závěr <b>tučně</b> & <font color="red">x</font>' in text  # printed literally, never as markup
    assert "MUDr. <i>Test</i>" in text
    assert "635722/1013" not in text
    assert f"RČ {REDACTED_BIRTH_NUMBER} konec" in text


def test_a_conclusion_longer_than_a_page_flows_onto_the_next_pages(story_handover: Handover) -> None:
    long = SignedConclusion(
        text="\n".join(f"{i}. řádek závěru" for i in range(1, 301)),
        signer_display="MUDr. Lenka Marková",
        signed_at=SIGNED_AT,
    )
    context = build_context(story_handover, GP, EXPORTED_AT, conclusion=long)

    text = _text(render_mdt_record_pdf(story_handover, GP, EXPORTED_AT, conclusion=long))

    assert "1. řádek závěru" in text
    assert "300. řádek závěru" in text
    assert _norm(context["footer"]) in text


def test_an_incomplete_layout_is_refused(story_handover: Handover, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(mdt_record_pdf._EndMarker, "draw", lambda self: None)

    with pytest.raises(MdtRecordError):
        render_mdt_record_pdf(story_handover, GP, EXPORTED_AT)
